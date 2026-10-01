# Bộ file Excel GIẢ cùng cấu trúc hệ thống (không dùng dữ liệu thật)
import openpyxl, random, sys, os
random.seed(7)
N = int(sys.argv[1]) if len(sys.argv)>1 else 25000
OUT = sys.argv[2] if len(sys.argv)>2 else 'gia'
os.makedirs(OUT, exist_ok=True)
XA = [('540034','Phường Gia Lộc',[('TXN0543404','Gia Lộc 2','07',['54003401','54003402']),('TXN0543405','Gia Lộc','23',['54003403','54003407'])]),
      ('540035','Phường Gò Dầu',[('TXN0543502','Phường Gò Dầu','05',['54003508','54003511']),('TXN0543501','Gia Bình','12',['54003505','54003501'])]),
      ('540083','Xã Truông Mít',[('TXN0548302','Truông Mít','25',['54008306','54008309'])])]
CT = [('01','HONGHEO','Cho vay ưu đãi hộ nghèo'),('02','HSSV','Cho vay học sinh, sinh viên có hoàn cảnh khó khăn'),('03','GQVL','Cho vay giải quyết việc làm'),('06','NSVSMT','Cho vay nước sạch và vệ sinh môi trường nông thôn'),('19','HCN_QD15','Cho vay hộ cận nghèo theo QĐ 15')]
HO = ['Nguyễn','Trần','Lê','Phạm','Võ','Huỳnh','Đặng','Bùi']; DEM=['Văn','Thị','Ngọc','Minh','Thanh']; TEN=['An','Bình','Cường','Dung','Giả','Hạnh','Khoa','Lan','Mai','Nam','Oanh','Phúc','Quý','Sang','Tâm','Uyên','Vy','Xuân','Yến','Thử']
thon_info = {}
for xa,tx,ds in XA:
    for md,td,ng,th in ds:
        for t in th: thon_info[t]=(xa,tx,md,td,ng)
to_list=[]
for i,t in enumerate(thon_info):
    for j in range(6): to_list.append(('%07d'%(26000+i*10+j), t, random.choice(['11','12','13','14'])))
def ten(): return random.choice(HO)+' '+random.choice(DEM)+' '+random.choice(TEN)
khs=[]; ku_n=0; rows=[]
nkh = int(N*0.75)
for k in range(nkh):
    to,thon,dv = random.choice(to_list); makh='48%08d'%k
    kh=(makh,ten(),'%02d/%02d/%d'%(random.randint(1,28),random.randint(1,12),random.randint(1950,2000)),'0721%08d'%random.randint(0,10**8-1),'%02d/%02d/2021'%(random.randint(1,28),random.randint(1,12)),'CỤC CS QLHC',' Ấp giả - xã giả - Tây Ninh',to,thon,dv)
    khs.append(kh)
while len(rows)<N:
    kh=random.choice(khs); ct=random.choice(CT); ku_n+=1
    dn=random.choice([10,20,30,40,50,100])*1_000_000; tt=random.random()
    qh = dn if tt<0.003 else 0; kn = dn if 0.003<=tt<0.005 else 0; th=dn-qh-kn
    rows.append(dict(kh=kh,ku='66000007%08d'%ku_n,ct=ct,gn=dn+random.choice([0,5,10])*1_000_000,dn=dn,th=th,qh=qh,kn=kn,nv='%02d/%02d/2023'%(random.randint(1,28),random.randint(1,12)),ndt=random.choice(['','','INV2706190050215']),nguon=random.choice(['1','1','1','2']), t105=random.randint(0,9)*1_000_000))
def sheet(ws, tieu, cot, data, them_nhieu=True):
    ws.cell(2,2,tieu)
    for j,c in enumerate(cot): ws.cell(5,j+2,c)
    r=6
    for i,d in enumerate(data):
        if them_nhieu and i==3: r+=1  # dòng trống
        if them_nhieu and i==len(data)//2:
            for j,c in enumerate(cot): ws.cell(r,j+2,c)   # tiêu đề lặp (ngắt trang)
            r+=1
        for j,v in enumerate(d): ws.cell(r,j+2,v)
        r+=1
    if them_nhieu:
        ws.cell(r,2,'Tổng cộng'); r+=2; ws.cell(r,2,'Người lập biểu'); ws.cell(r,6,'Kiểm soát')
def luu(wb,ten): wb.save(os.path.join(OUT,ten))
# 1 HSTD
C1=['Mã PGD','Mã xã','Tên xã','Mã thôn','Tên thôn','ĐVUT','Mã tổ','Tên tổ trưởng','Mã khách hàng','Tên khách hàng','Ngày sinh','CCCD','Ngày cấp','Nơi cấp','Địa chỉ','Mã món vay','Chương trình','Tên chương trình','Số tiền giải ngân','Tổng dư nợ','Dư nợ trong hạn','Dư nợ quá hạn','Dư nợ khoanh','Ngày vay','Ngày giải ngân đầu tiên','Số tiền gia hạn','Số tháng gia hạn','Tổng sô tháng gia hạn','Nguồn vốn','Lưu vị trong năm','Lãi tồn trong hạn','Lãi tồn quá hạn','Lãi tồn AH','Lãi tồn','Ngày nhập học','Ngày ra trường','HSSV','Số TK','Tên TK','105 đầu tháng','105 Ngày BC','Mã NĐT']
seen=set(); d1=[]
for x in rows:
    k=x['kh']; xa=thon_info[k[8]]; first = k[0] not in seen; seen.add(k[0])
    d1.append(['004820',xa[0],xa[1],k[8],'Thôn '+k[8][-2:],k[9],k[7],'TỔ TRƯỞNG '+k[7][-3:],k[0],k[1],k[2],k[3],k[4],k[5],k[6],x['ku'],x['ct'][0],x['ct'][2],x['gn'],x['dn'],x['th'],x['qh'],x['kn'],x['nv'],x['nv'],0,0,0,x['nguon'],'1',0,0,0,0,'01/09/2021' if x['ct'][0]=='02' else None,'01/06/2025' if x['ct'][0]=='02' else None,x['ct'][0],'48%08d'%int(k[0][2:]) if first else '', k[1].upper() if first else '', x['t105'] if first else 0, x['t105'] if first else 0, x['ndt']])
wb=openpyxl.Workbook(); ws=wb.active; ws.title='BCQUERY'; sheet(ws,'10. Sao kê chi tiết (DL Ngày)',C1,d1); luu(wb,'Ho so tin dung chi tiet 31-08-2026.xlsx')
# 6 TDN (gộp theo thôn × CT × nguồn × NĐT)
agg={}
for x in rows:
    k=x['kh']; key=(k[8],x['ct'][0],x['nguon'],x['ndt'])
    a=agg.setdefault(key,[set(),0,0,0,0]); a[0].add(k[0]); a[1]+=x['dn']; a[2]+=x['th']; a[3]+=x['qh']; a[4]+=x['kn']
C6=['Mã PGD','Nguồn vốn','Mã xã','Tên xã','Mã điểm giao dịch','Điểm giao dịch','Mã Thôn','Tên Thôn','Chương Trình','Chương trình VT','Tên Chương trình','Số KH','Tổng Dư nợ','DN Trong hạn','DN Quá hạn','DN Khoanh','Mã NĐT','Ngày báo cáo']
ctm={c[0]:c for c in CT}; d6=[]
for (t,c,ng,nd),a in sorted(agg.items()):
    xa=thon_info[t]; d6.append(['004820',ng,xa[0],xa[1],xa[2],xa[3],t,'Thôn '+t[-2:],c,ctm[c][1],ctm[c][2],len(a[0]),a[1],a[2],a[3],a[4],nd,'31/08/2026'])
wb=openpyxl.Workbook(); ws=wb.active; ws.title='BCQUERY'; sheet(ws,'1. TỔNG DƯ NỢ THEO CHƯƠNG TRÌNH',C6,d6,False); luu(wb,'Tong_du_no_2026-08-31.xlsx')
# 4 NQH (có Ngày báo cáo)
C4=['Mã PGD','Mã xã','Tên xã','Thôn','Mã điểm GD','Tên điểm GD','Mã tổ','Tên tổ trưởng','ĐVUT','Tên ĐVUT','Mã khách hàng','Tên khách hàng','Số khế ước','Dư nợ quá hạn','Chuyển QH trong tháng','Ngày chuyển quá hạn','Nguồn vốn','Chương trình','Chương trình VT','Tên Chương trình','Số dư 105','Ngày báo cáo']
d4=[]
for x in rows:
    if x['qh']>0:
        k=x['kh']; xa=thon_info[k[8]]; d4.append(['004820',xa[0],xa[1],k[8],xa[2],xa[3],k[7],'TỔ TRƯỞNG',k[9],'Hội',k[0],k[1],x['ku'],x['qh'],x['qh'] if random.random()<0.2 else 0,'23/08/2026','1',x['ct'][0],x['ct'][1],x['ct'][2],0,'31/08/2026'])
wb=openpyxl.Workbook(); ws=wb.active; ws.title='BCQUERY'; sheet(ws,'2. SAO KÊ DANH SÁCH NỢ QUÁ HẠN',C4,d4); luu(wb,'No qua han  31-08-2026.xlsx')
# 5 NK (không có ngày trong file) — tên file đảo từ, không dấu
C5=['Mã PGD','Ngày GDXA','Điểm GDXA','Tên điểm GDXA','Mã xã','Tên xã','Mã khách hàng','Tên khách hàng','Mã tổ','Tên tổ trưởng','Số khế ước','ĐVUT','Dư nợ khoanh','Ngày hiệu lực','Ngày hết hạn khoanh','Nguồn vốn','Chương trình','Tên chương trình','Nguyên nhân']
d5=[]
for x in rows:
    if x['kn']>0:
        k=x['kh']; xa=thon_info[k[8]]; d5.append(['004820',xa[4],xa[2],xa[3],xa[0],xa[1],k[0],k[1],k[7],'TT',x['ku'],k[9],x['kn'],'15/12/2022','15/12/2027','1',x['ct'][0],x['ct'][2],'Vắng mặt'])
wb=openpyxl.Workbook(); ws=wb.active; ws.title='BCQUERY'; sheet(ws,'13. Danh sách nợ khoanh (DL Tháng)',C5,d5); luu(wb,'Khoanh No_2026-08-31.xlsx')
# 3 KHĐ + sheet phụ của anh
C3=['Mã CN','Mã PGD','Mã thôn/ấp','Mã điểm GDX','Tên điểm GDX','ĐVUT','Mã tổ','Tên tổ trưởng','Mã khách hàng','Tên khách hàng','Số khế ước','Mã sản phẩm','Ngày đăng ký khoản vay','Ngày giải ngân đầu tiên','Ngày đến hạn gốc','Ngày đến hạn GH','Ngày đến hạn GDXA','Chương trình','Tổng dư nợ','Dư nợ trong hạn','Dư nợ quá hạn','Dư nợ khoanh','Ngày giao dịch gần nhất','Lãi đã thu','Lãi tồn']
d3=[]
for x in random.sample(rows, max(5,N//80)):
    k=x['kh']; xa=thon_info[k[8]]; d3.append(['005420','004820',k[8],xa[2],xa[3],k[9],k[7],'TT',k[0],k[1],x['ku'],'031MC1','03/04/2024','07/04/2024','07/04/2029','07/04/2029','07/04/2029',x['ct'][0],x['dn'],x['th'],x['qh'],x['kn'],'07/05/2026',1000000,200000])
wb=openpyxl.Workbook(); ws=wb.active; ws.title='BCQUERY'; sheet(ws,'14. Sao kê món vay N tháng không hoạt động (DL Tháng)',C3,d3)
w2=wb.create_sheet('Sheet4'); w2.append(['MACT','TENCT','TENVT']); w2.append(['01','Cho vay ưu đãi hộ nghèo','HONGHEO'])
w3=wb.create_sheet('BC Gia Lộc'); w3.append(['CHI NHÁNH']); w3.append(['DANH SÁCH MÓN VAY 3 THÁNG TRỞ LÊN KHÔNG HOẠT ĐỘNG']); w3.append(['ĐẾN NGÀY 31/08/2026'])
luu(wb,'Mon_vay_03_thang_KHD_2026-08-31.xlsx')
# 7 TT — mã tổ để dạng SỐ (mất số 0 đầu) để thử bù
C7=['Mã PGD','Mã xã','Tên xã','Mã điểm GDXA','Tên điểm GDXA','Ngày GDXA','Mã thôn','Tên thôn','ĐVUT','Mã tổ trưởng','Tên tổ trưởng','Năm sinh tổ trưởng','SĐT Tổ trưởng','Mã KH tổ phó','Tên Tổ phó','Trạng thái','Loại tổ']
d7=[]
for to,thon,dv in to_list:
    xa=thon_info[thon]; d7.append(['004820',xa[0],xa[1],xa[2],xa[3]+' ',xa[4],thon,'Thôn '+thon[-2:],dv,int(to),'TỔ TRƯỞNG '+to[-3:],'01/01/1970','09%08d'%int(to),'7000000001','TỔ PHÓ','A','01'])
wb=openpyxl.Workbook(); ws=wb.active; ws.title='BCQUERY'; sheet(ws,'12. THÔNG TIN TỔ TRƯỞNG',C7,d7); luu(wb,'Thong tin to truong T8 2026.xlsx')
print('món',len(rows),'KH',len(seen),'NQH',len(d4),'NK',len(d5),'KHĐ',len(d3),'TDN',len(d6),'tổ',len(d7))
print('TONG dn',sum(x['dn'] for x in rows),'qh',sum(x['qh'] for x in rows),'kn',sum(x['kn'] for x in rows))

# ---------- 3.86: Mẫu 31 "Tạo hồ sơ tín dụng chi tiết theo kỳ số liệu" (175 cột) — GIẢ, thay Mẫu 10 khi gọi: taogia.py N thư_mục m31 ----------
C31 = ['Mã CN', 'Mã PGD', 'Tên PGD', 'Mã xã', 'Tên xã', 'Mã thôn', 'Tên thôn', 'Mã KH', 'Tên KH', 'Ngày sinh', 'Phân loại', 'Loại KH', 'Giới tính', 'Mã dân tộc', 'Tên DT', 'Số CMND', 'Nơi cấp CMND', 'Ngày cấp CMND', 'Tên vợ/chồng', 'Địa chỉ', 'Số điện thoại', 'Mã tổ', 'Loại tổ', 'Mã CIF TT', 'Tên tổ', 'Mã ĐVUT', 'Tên ĐVUT', 'Số khế ước', 'Ngày vay', 'Ngày ĐH theo hợp đồng', 'Ngày ĐH theo Gia hạn', 'Ngày ĐH theo GDXA', 'Thời hạn vay', 'Lãi suất', 'Ngày B đầu trả gốc', 'Ngày B đầu trả lãi', 'Ngày KTAHSV', 'Ngày GDXA', 'Hình thức vay', 'Tình trạng món vay', 'Cấp QL vốn', 'Tên cấp QLV', 'Đối tượng thụ hưởng', 'Tên ĐTTH', 'Mã sản phẩm cụ thể', 'Mã SP tín dụng', 'Mã DM sản phẩm', 'Nguồn vốn', 'Mã chương trình', 'Tên chương trình', 'Mã Quyết định', 'Tên Quyết định', 'Mức vay', 'Tổng giải ngân', 'Ngày GN đầu tiên', 'Ngày GN cuối cùng', 'Dư nợ trong hạn', 'Dư nợ quá hạn', 'Dư nợ khoanh', 'Tổng dư nợ', 'Gốc đến hạn LK', 'Gốc đã trả', 'Gốc xóa', 'Lãi xóa', 'Lũy kế thu lãi', 'Lũy kế thu lãi TH', 'Lãi tồn TH', 'Lũy kế thu lãi QH', 'Lãi tồn QH', 'Lãi DT chưa đến hạn', 'Lãi DT trong tháng', 'Lãi TT trong tháng', 'Thu lãi TH tháng', 'Thu lãi QH Tháng', 'Giải ngân trong tháng', 'Đảo khoản GN tháng', 'Lưu vụ trong tháng', 'Gia hạn trong tháng', 'Chuyển QH trong tháng', 'Chuyển khoanh trong tháng', 'Thu nợ TH tháng', 'Thu nợ QH tháng', 'Thu nợ khoanh tháng', 'Gốc xóa trong tháng', 'Lãi DT Quý', 'Lãi TT Quý', 'Thu lãi TH Quý', 'Thu lãi QH Quý', 'Giải ngân trong Quý', 'Đảo khoản GN Quý', 'Lưu vụ trong Quý', 'Gia hạn trong Quý', 'CQH trong Quý', 'Chuyển Khoanh Quý', 'Thu nợ TH Quý', 'Thu nợ QH Quý', 'Thu nợ Khoanh Quý', 'Gốc xóa trong Quý', 'Lãi DT Năm', 'Lãi TT Năm', 'Thu lãi TH Năm', 'Thu lãi QH Năm', 'Giải ngân Năm', 'Đảo khoản GN Năm', 'Lưu vụ Năm', 'Gia hạn Năm', 'CQH Năm', 'Chuyển Khoanh Năm', 'Thu nợ TH Năm', 'Thu nợ QH Năm', 'Thu nợ Khoanh Năm', 'Xóa trong Năm', 'Tài khoản TH', 'Tài khoản QH', 'Tài khoản khoanh', 'Tài khoản thu lãi', 'Tổng thu nợ TH', 'Tổng thu nợ QH', 'Tổng thu nợ Khoanh', 'Tổng gia hạn nợ', 'Số lần đã gia hạn', 'Ngày GH gần nhất', 'Số tháng đã GH', 'Tổng chuyển nợ QH', 'Ngày CNQH gần nhất', 'Tổng chuyển nợ khoanh', 'Ngày hết hạn Khoanh', 'Gốc hết hạn Khoanh', 'Ngày lưu vụ', 'Tồn RPA', 'Ngày giao dịch gần nhất', 'Ngày dự thu', 'Tổng Lãi phải HT', 'Lãi chưa hỗ trợ', 'Chuẩn nghèo ĐP', 'Mã nhà đầu tư', 'Mã PNKT51', 'Tên PNKT51', 'Mã PNKT52', 'Mã HQĐT', 'Tên HQĐT', 'Giá trị HQĐT1', 'Giá trị HQĐT2', 'Ký quý', 'Mục đích nhà', 'Mục đích 30A', 'Mã dự án', 'Mã HSSV', 'Tên HSSV', 'CMND HSSV', 'Mã trường HSSV', 'Tên trường', 'Mã Hệ đào tạo', 'Tên Hệ ĐT', 'Mã ngành ĐT', 'Tên Ngành ĐT', 'Mã ĐT học phí', 'Tên ĐT học phí', 'Ngày nhập học', 'Ngày ra trường', 'Số ATM', 'Đơn vị cấp thẻ', 'Tên LĐXK', 'Ngày sinh LĐXK', 'Số thẻ LĐXK', 'Ngày hợp đồng XK', 'Ngày hết hạn XK', 'Lĩnh vực LĐXK', 'Công ty xuất khẩu', 'Quốc gia xuất khẩu', 'Tên Quốc Gia', 'Công ty môi giới XK', 'Sổ tiết kiệm 105', 'Số dư tiền gửi 105', 'Ngày số liệu']
def mau31(ten_file, ky_ngay, lui):
    """lui=True: tạo kỳ TRƯỚC (dư nợ đầu kỳ) để thử công thức dư nợ tháng trước + phát sinh = dư nợ tháng này"""
    ix = {c:i for i,c in enumerate(C31)}
    out = []
    def dong(vals):
        r = ['']*len(C31)
        for c,v in vals.items(): r[ix[c]] = v
        return r
    so_tk = {}
    for n,x in enumerate(rows):
        k=x['kh']; xa=thon_info[k[8]]
        gn_t = x['dn'] if n%200==0 else 0          # món giải ngân trong tháng
        thu = 0 if gn_t else (1_000_000 if n%7==0 else 0)
        dn = x['dn'] - (0 if not lui else (gn_t - thu))   # kỳ trước: dn + thu − gn
        if lui and gn_t: continue                    # món mới giải ngân tháng này chưa có ở kỳ trước
        v = {'Mã CN':'005420','Mã PGD':'004820','Tên PGD':'PGD GIẢ','Mã xã':xa[0],'Tên xã':xa[1],'Mã thôn':k[8],'Tên thôn':'Thôn '+k[8][-2:],
             'Mã KH':k[0],'Tên KH':k[1],'Ngày sinh':k[2],'Phân loại':'E','Loại KH':'1','Giới tính':'01','Tên DT':'Kinh','Số CMND':k[3],'Nơi cấp CMND':k[5],'Ngày cấp CMND':k[4],
             'Tên vợ/chồng':'VỢ CHỒNG GIẢ','Địa chỉ':k[6],'Số điện thoại':'09%08d'%int(k[0][2:]),'Mã tổ':k[7],'Loại tổ':'01','Tên tổ':'Tổ '+k[7][-3:],'Mã ĐVUT':k[9],
             'Số khế ước':x['ku'],'Ngày vay':x['nv'],'Ngày ĐH theo hợp đồng':'20/03/2029','Ngày ĐH theo Gia hạn':'20/03/2029','Ngày ĐH theo GDXA':'19/04/2029','Thời hạn vay':'Trung hạn',
             'Lãi suất':8.4,'Ngày GDXA':'19','Tình trạng món vay':'OPEN','Nguồn vốn':x['nguon'],'Mã chương trình':x['ct'][0],'Tên chương trình':x['ct'][2],'Mức vay':float(x['gn']),'Tổng giải ngân':float(x['gn']),
             'Dư nợ trong hạn':float(dn-x['qh']-x['kn']),'Dư nợ quá hạn':float(x['qh']),'Dư nợ khoanh':float(x['kn']),'Tổng dư nợ':float(dn),
             'Giải ngân trong tháng':0.0 if lui else float(gn_t),'Thu nợ TH tháng':0.0 if lui else float(thu),'Đảo khoản GN tháng':0.0,'Thu nợ QH tháng':0.0,'Thu nợ khoanh tháng':0.0,'Gốc xóa trong tháng':0.0,
             'Lãi tồn TH':0.0,'Lãi tồn QH':0.0,'Ngày giao dịch gần nhất':'19/06/2026','Mã nhà đầu tư':x['ndt'],
             'Sổ tiết kiệm 105':'', 'Số dư tiền gửi 105':0.0, 'Ngày số liệu':ky_ngay}
        if k[0] not in so_tk: so_tk[k[0]]=1; v['Sổ tiết kiệm 105']='48%08d'%int(k[0][2:]); v['Số dư tiền gửi 105']=float(x['t105'])
        if x['ct'][0]=='02': v.update({'Tên HSSV':'SINH VIÊN GIẢ','Tên trường':'TRƯỜNG GIẢ','Ngày nhập học':'01/09/2021','Ngày ra trường':'01/06/2025'})
        out.append(dong(v))
        if n%1000==5:   # khách có 2 sổ 105 → hệ thống xuất 2 dòng cùng khế ước
            v2=dict(v); v2['Sổ tiết kiệm 105']='49%08d'%int(k[0][2:]); v2['Số dư tiền gửi 105']=500000.0; out.append(dong(v2))
    # món XKLĐ, món đã tất toán, khách chỉ gửi tiết kiệm
    k=khs[0]; xa=thon_info[k[8]]
    base={'Mã CN':'005420','Mã PGD':'004820','Mã xã':xa[0],'Tên xã':xa[1],'Mã thôn':k[8],'Mã KH':k[0],'Tên KH':k[1],'Ngày sinh':k[2],'Số CMND':k[3],'Mã tổ':k[7],'Mã ĐVUT':k[9],'Ngày số liệu':ky_ngay,'Lãi suất':0.0,'Mức vay':0.0,'Tổng giải ngân':0.0,'Tổng dư nợ':0.0,'Dư nợ trong hạn':0.0,'Dư nợ quá hạn':0.0,'Dư nợ khoanh':0.0,'Số dư tiền gửi 105':0.0}
    if not lui:
        out.append(dong(dict(base, **{'Số khế ước':'6600009900000001','Tình trạng món vay':'OPEN','Mã chương trình':'04','Tên chương trình':'Cho vay ĐTCS đi lao động có thời hạn ở nước ngoài','Tổng giải ngân':100000000.0,'Tổng dư nợ':100000000.0,'Dư nợ trong hạn':100000000.0,'Giải ngân trong tháng':100000000.0,
            'Tên LĐXK':'NGƯỜI LAO ĐỘNG GIẢ','Quốc gia xuất khẩu':'JP','Tên Quốc Gia':'Nhật Bản','Công ty xuất khẩu':'CÔNG TY GIẢ'})))
    out.append(dong(dict(base, **{'Số khế ước':'6600009900000002','Tình trạng món vay':'CLOSE','Mã chương trình':'03','Thu nợ TH Năm':20000000.0})))
    for j in range(30):
        kk=khs[-1-j]; out.append(dong({'Mã CN':'005420','Mã PGD':'004820','Mã thôn':'48200000','Mã KH':'47%08d'%j,'Tên KH':'KHÁCH GỬI TK GIẢ %d'%j,'Phân loại':'1','Sổ tiết kiệm 105':'47%08d'%j,'Số dư tiền gửi 105':1000000.0,'Ngày số liệu':ky_ngay,
            'Tổng dư nợ':0.0,'Dư nợ trong hạn':0.0,'Dư nợ quá hạn':0.0,'Dư nợ khoanh':0.0}))
    wb=openpyxl.Workbook(); ws=wb.active; ws.title='BCQUERY'
    ws.cell(2,2,'31 - Tạo hồ sơ tín dụng chi tiết theo kỳ số liệu')
    for j,c in enumerate(C31): ws.cell(5,j+2,c)
    for i,r in enumerate(out):
        for j,v in enumerate(r):
            if v!='': ws.cell(6+i,j+2,v)
    luu(wb,ten_file)
    return len(out)
if len(sys.argv)>3 and sys.argv[3]=='m31':
    os.remove(os.path.join(OUT,'Ho so tin dung chi tiet 31-08-2026.xlsx'))
    print('Mẫu 31 T8:', mau31('Ho_so_tin_dung_chi_tiet_31-08-2026.XLSX','31/08/2026',False), 'dòng · T7:', mau31('Ho so tin dung chi tiet 31-07-2026 mau31.xlsx','31/07/2026',True), 'dòng')
