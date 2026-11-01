#!/usr/bin/env python3
"""Dựng bảng ĐÁP ÁN số từ so_lieu_tong_hop.json của giao ban (docs/GIAO_BAN_KY_NANG.md mục 7).
Chạy: python3 tools/dap_an_giao_ban.py <Diem> <json kỳ 1> [<json kỳ 2> ...]  -> DAP_AN_<Diem>.xlsx (hiện lưới) + DAP_AN_<Diem>.csv
File json là DỮ LIỆU THẬT (Drive 00_PHOI_HOP/BAN_GIAO_APP) — chạy ngoài repo, KHÔNG commit đầu ra.
CSV dạng dài: ky, diem, hoi (HND|HPN|HCCB|DTN|TT|CONG|trống), chi_tieu, gia_tri — app xuất cùng dạng để so bằng máy."""
import json, sys, csv, openpyxl
from openpyxl.styles import Font
DIEM = sys.argv[1]; files = sys.argv[2:]
HOI = ["HND", "HPN", "HCCB", "DTN"]
CT_HOI = ["so_to", "so_kh_du_no", "so_kh_tat_no_con_tg", "so_ho", "du_no", "du_no_truoc", "du_no_tang", "nqh", "khoanh",
          "cho_vay", "thu_no", "tien_gui", "tien_gui_tang", "lai_ton", "lai_ton_tang", "khd_mon", "khd_du", "ndh", "ndh_ho"]
wb = openpyxl.Workbook(); wb.remove(wb.active)
rows_csv = []
def r3(v): return round(v, 3) if isinstance(v, float) else v
for f in files:
    D = json.load(open(f)); d = D["diem"][DIEM]; ky = D["ngay_sl"]
    ws = wb.create_sheet(("SL" + ky.replace("/", "-"))[:31]); ws.sheet_view.showGridLines = True
    ws.append(["Kỳ số liệu " + ky + " — so với " + D.get("ngay_truoc", "") + " — điểm " + DIEM]); ws["A1"].font = Font(bold=True)
    ws.append(["Chỉ tiêu (khóa json)"] + HOI + ["Trực tiếp", "Cộng"]); [setattr(c, "font", Font(bold=True)) for c in ws[2]]
    tt = d.get("truc_tiep", {})
    for k in CT_HOI:
        vs = [d["hoi"][h].get(k) for h in HOI]
        t = tt.get({"so_ho": "so_ho"}.get(k, k)) if k in ("du_no", "nqh", "khoanh", "cho_vay", "thu_no", "lai_ton", "so_ho", "du_no_tang") else None
        cong = sum(v for v in vs if isinstance(v, (int, float))) + (t or 0) if all(isinstance(v, (int, float)) for v in vs) else None
        ws.append([k] + [r3(v) for v in vs] + [r3(t), r3(cong)])
        for h, v in zip(HOI + ["TT", "CONG"], vs + [t, cong]): rows_csv.append([ky, DIEM, h, k, r3(v)])
    ws.append([]); ws.append(["Danh sách (số dòng / tổng tiền, triệu đồng)"]); ws.cell(ws.max_row, 1).font = Font(bold=True)
    def tong(lst, k="so_tien"): return round(sum(x.get(k) or 0 for x in lst), 3)
    ds = [("nqh_list", "Nợ quá hạn", "so_tien"), ("khoanh_list", "Nợ khoanh", "so_tien"), ("khd_list", "Món 3 tháng KHĐ", "du_no"),
          ("gn_list", "Giải ngân trong tháng (mọi lần)", "so_tien"), ("ndh_list", "Nợ đến hạn tháng " + str(d.get("ndh_thang")), "so_tien"),
          ("ndh_sau_list", "Nợ đến hạn tháng sau", "so_tien"), ("ndh_kq", "Kết quả nợ đến hạn tháng trước", "so_tien")]
    for k, ten, tk in ds:
        lst = d.get(k) or []
        ws.append([ten + " (" + k + ")", len(lst), tong(lst, tk)]); rows_csv += [[ky, DIEM, "", k + ".so_dong", len(lst)], [ky, DIEM, "", k + ".tong", tong(lst, tk)]]
    gn1 = [x for x in d.get("gn_list") or [] if x.get("ky_dau")]
    ws.append(["Khoản vay mới = giải ngân lần 1 (gn_list.ky_dau)", len(gn1), tong(gn1)]); rows_csv += [[ky, DIEM, "", "gn_lan1.so_dong", len(gn1)], [ky, DIEM, "", "gn_lan1.tong", tong(gn1)]]
    ws.append(["Số tổ (to)", len(d.get("to") or [])]); rows_csv.append([ky, DIEM, "", "to.so_dong", len(d.get("to") or [])])
    if d.get("moc"): ws.append([]); ws.append(["Mốc " + str(d["moc"].get("ngay"))] + [str(k) + "=" + str(r3(v)) for k, v in d["moc"].items() if k != "ngay"])
    # bảng tổ
    wt = wb.create_sheet(("To_" + ky.replace("/", "-"))[:31]); wt.sheet_view.showGridLines = True
    kt = ["ma", "ten", "hoi", "thon", "so_kh", "kh_du_no", "du_no", "nqh", "khoanh", "cho_vay", "thu_no", "tien_gui", "lai_ton", "lt3_mon", "lt3_tien", "tk_thap"]
    wt.append(kt); [setattr(c, "font", Font(bold=True)) for c in wt[1]]
    for t in d.get("to") or []: wt.append([r3(t.get(k)) for k in kt])
    for s in (ws, wt):
        s.column_dimensions["A"].width = 44 if s is ws else 10
        if s is wt: s.column_dimensions["B"].width = 24
wb.save("DAP_AN_%s.xlsx" % DIEM)
with open("DAP_AN_%s.csv" % DIEM, "w", newline="", encoding="utf-8-sig") as fo:
    w = csv.writer(fo); w.writerow(["ky", "diem", "hoi", "chi_tieu", "gia_tri"]); w.writerows(rows_csv)
print("ok", len(rows_csv))
